import os
import re
import shutil
from datetime import datetime
from typing import Optional
from openpyxl import load_workbook
from config.settings import INPUT_EXCEL_PATH, OUTPUT_DIR, SCENARIO_NAME

SCENARIO_SHEET = "Scenario"
SCENARIO_OUTPUT_COLUMNS = ["Status", "Screenshot"]
DATA_OUTPUT_COLUMNS = ["Status", "Actual"]

# Matches: Scenario1_20260710_143210.xlsx
_FILENAME_RE = re.compile(r"^(?P<scenario>.+)_(?P<date>\d{8})_(?P<time>\d{6})\.xlsx$")


def _today_existing_file(scenario_name: str) -> Optional[str]:
    if not os.path.isdir(OUTPUT_DIR):
        return None
    today = datetime.now().strftime("%Y%m%d")
    for fname in os.listdir(OUTPUT_DIR):
        match = _FILENAME_RE.match(fname)
        if match and match.group("scenario") == scenario_name and match.group("date") == today:
            return os.path.join(OUTPUT_DIR, fname)
    return None


def _add_output_columns(ws, columns: list[str]) -> None:
    headers = [c.value for c in ws[1]]
    for col in columns:
        if col not in headers:
            ws.cell(row=1, column=ws.max_column + 1, value=col)


def _build_output_from_input(output_path: str, input_path: str = INPUT_EXCEL_PATH) -> None:
    shutil.copy(input_path, output_path)
    wb = load_workbook(output_path)

    _add_output_columns(wb[SCENARIO_SHEET], SCENARIO_OUTPUT_COLUMNS)

    for sheet_name in wb.sheetnames:
        if sheet_name == SCENARIO_SHEET:
            continue
        _add_output_columns(wb[sheet_name], DATA_OUTPUT_COLUMNS)

    wb.save(output_path)


def get_output_path(scenario_name: str = SCENARIO_NAME, input_path: str | None = None) -> str:

    os.makedirs(OUTPUT_DIR, exist_ok=True)

    existing = _today_existing_file(scenario_name)
    if existing:
        os.remove(existing)  # drop the stale snapshot, don't just rename it

    now = datetime.now()
    new_name = f"{scenario_name}_{now.strftime('%Y%m%d')}_{now.strftime('%H%M%S')}.xlsx"
    new_path = os.path.join(OUTPUT_DIR, new_name)

    _build_output_from_input(
        new_path,
        input_path or INPUT_EXCEL_PATH,
    )  # always fresh from the selected input file
    return new_path
