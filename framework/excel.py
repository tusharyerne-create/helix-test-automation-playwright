from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from config.settings import INPUT_EXCEL_PATH


class ExcelManager:
    def __init__(self, path: str = INPUT_EXCEL_PATH):
        self.path = path
        self.workbook = load_workbook(path, data_only=False)

    def get_sheet(self, sheet_name: str) -> Worksheet:
        if sheet_name not in self.workbook.sheetnames:
            raise KeyError(f"Sheet '{sheet_name}' not found in {self.path}")
        return self.workbook[sheet_name]

    def read_sheet_as_dicts(self, sheet_name: str) -> list[dict]:

        ws = self.get_sheet(sheet_name)
        headers = [c.value for c in ws[1]]
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if all(v is None for v in row):
                continue
            rows.append({headers[i]: row[i] for i in range(len(headers))})
        return rows

    def find_column(self, sheet_name: str, header_name: str) -> int | None:
        ws = self.get_sheet(sheet_name)
        for idx, cell in enumerate(ws[1], start=1):
            if cell.value == header_name:
                return idx
        return None

    # ---------- writing ----------
    def write_cell(self, sheet_name: str, row: int, column: int, value) -> None:
      
        self.get_sheet(sheet_name).cell(row=row, column=column).value = value

    def write_cell_by_header(self, sheet_name: str, row: int, header_name: str, value) -> None:
        col = self.find_column(sheet_name, header_name)
        if col is None:
            raise KeyError(f"Header '{header_name}' not found in sheet '{sheet_name}'")
        self.write_cell(sheet_name, row, col, value)

    def save(self) -> None:
        self.workbook.save(self.path)